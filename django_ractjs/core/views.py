from .models import Files
from rest_framework import viewsets
from rest_framework.decorators import api_view
from .serializers import FilesSerializer
import csv
from django.http import JsonResponse
from django.core.files.storage import default_storage
import json
import csv
import pandas as pd
from django.db import models
import numpy as np
from django.views.decorators.csrf import csrf_exempt
import pypyodbc as odbc
from django.shortcuts import render, HttpResponse
import re
import math
from sklearn.preprocessing import LabelEncoder
import openpyxl
from django.contrib.auth import authenticate, login
from django.shortcuts import render, redirect
from django.contrib.auth import get_user_model
from django.contrib.auth import logout

DRIVER = "SQL Server"
SERVER_NAME = "LAPTOP-H3TEL2C9\SQLEXPRESS"
DATABASE_NAME = "database1"


def read_file(file_path):
    # get the file extension
    file_ext = file_path.split('.')[-1]

    if file_ext == 'csv':
        # read the CSV file into a pandas DataFrame
        df = pd.read_csv(file_path)
    elif file_ext == 'xlsx':
        # read the Excel file into a pandas DataFrame
        workbook = openpyxl.load_workbook(file_path)
        sheet = workbook.active
        data = sheet.values
        headers = next(data)
        df = pd.DataFrame(data, columns=headers)
    else:
        raise ValueError('Unsupported file format')

    return df


def connection_string(driver, server_name, db_name):
    conn_string = f"""
        DRIVER={{{driver}}};
        SERVER={server_name};
        DATABASE={db_name};
        Trust_Connection=yes;
    """
    return conn_string


try:
    conn = odbc.connect(connection_string(DRIVER, SERVER_NAME, DATABASE_NAME))
except odbc.DatabaseError as e:
    print("Database Error")
    print(str(e.value[1]))
except odbc.Error as e:
    print("Connection Error")
    print(str(e.value[1]))


def select_table_query(tableName):
    sql_select_all = f"""SELECT * FROM {tableName};"""
    return sql_select_all


def get_column_name(df):
    columns = []
    dataTypes = {}
    tableCols = []
    for columnName in df:
        tableColName = re.sub("[^0-9a-zA-Z]+", "_", columnName)
        tableColName = tableColName + tableColName[-1]
        columns.append(columnName)
        tableCols.append(tableColName)

        if df[columnName].dtype == "Int64":
            dataTypes[tableColName] = "NVARCHAR(MAX)"

        elif df[columnName].dtype == "string" or df[columnName].dtype == "boolean":
            dataTypes[tableColName] = "VARCHAR(MAX)"

        elif df[columnName].dtype == "datetime64[ns]":
            dataTypes[tableColName] = "DATETIME"

        elif df[columnName].dtype == "Float64":
            dataTypes[tableColName] = "FLOAT"

    return columns, dataTypes, tableCols


def get_all_tables():
    tableName = []
    cursor = conn.cursor()

    cursor.execute("SELECT table_name FROM information_schema.tables")
    for table_name in cursor:
        temp = str(table_name)
        temp = temp[2: len(temp) - 3]
        tableName.append(temp)
    cursor.commit()

    cursor.close()
    return tableName


def create_table_query(tableName, dataTypes):
    sql_createTable = f"""
    CREATE TABLE dbo.{tableName} ("""

    for key in dataTypes:
        sql_createTable += f"""
            {key} {dataTypes[key]},"""

    sql_createTable = sql_createTable[: len(sql_createTable) - 1]

    sql_createTable += f"""
    );"""

    return sql_createTable


def insert_sql_query(tableName, columns):
    sql_insert = f"""
        INSERT INTO {tableName}("""

    for columnName in columns:
        sql_insert += f"""{columnName}, """

    sql_insert = sql_insert[: len(sql_insert) - 2]
    sql_insert += f""") VALUES ("""

    for columnName in columns:
        sql_insert += "?, "

    sql_insert = sql_insert[: len(sql_insert) - 2]
    sql_insert += ");"

    return sql_insert


def get_non_null_cols_and_record(columns, singleRec):
    new_cols = []
    new_singleRec = []
    i = 0
    for eachCol in singleRec:
        if not pd.isna(eachCol):
            new_cols.append(columns[i])
            new_singleRec.append(eachCol)
        i += 1

    return new_cols, new_singleRec


def getTables(request):
    returnObj = {}
    tables = get_all_tables()
    returnObj["status"] = True
    returnObj["data"] = tables

    return HttpResponse(json.dumps(returnObj), content_type="application/json")


@api_view(['GET'])
def getTableData(request, tableName):

    jsonDataArray = []
    columns = []
    returnObj = {}
    sql_get_all_data = select_table_query(tableName)

    cursor = conn.cursor()
    cursor.execute(sql_get_all_data)
    sql_table_data = cursor.fetchall()

    data = pd.DataFrame(sql_table_data)
    data = data.to_dict("split")

    columns = [i[0] for i in cursor.description]
    data = data["data"]

    for row in data:
        jsonData = {}
        i = 0
        for eachData in row:
            jsonData[columns[i]] = eachData
            i += 1
        jsonDataArray.append(jsonData)

    returnObj["status"] = True
    returnObj["data"] = jsonDataArray

    cursor.close()

    return HttpResponse(json.dumps(returnObj), content_type="application/json")


def sort_dataframe(df, sortColumn):
    if sortColumn == "All":
        sorted_df = df.apply(lambda col: pd.to_numeric(
            col, errors="coerce")).fillna(df)
        sorted_df = sorted_df.sort_values(by=sorted_df.columns.tolist())
    else:
        sorted_df = df.sort_values(by=sortColumn)

    return sorted_df


def transform_dataframe(df, transformationOptions, sortColumn):

    df = df.apply(pd.to_numeric, errors='ignore')

    for options in transformationOptions:
        num = options[0]

        if num == "1":
            # Remove duplicate rows
            df = df.drop_duplicates()
        elif num == "2":

            # Replace missing values with the mean of each column
            for col in df.columns:
                if df[col].dtype == 'int64' or df[col].dtype == 'float64':
                    col_mean = df[col].mean()
                    df[col].fillna(col_mean, inplace=True)
        elif num == "3":
            # Convert string columns to uppercase
            str_cols = df.select_dtypes(include='object').columns
            df[str_cols] = df[str_cols].apply(lambda x: x.str.upper())
        elif num == "4":
            # Remove columns with all missing values
            df = df.dropna(how='all', axis=1)
        elif num == "5":
            # Convert a categorical column to numeric
            cat_cols = df.select_dtypes(include=['object']).columns
            for col in cat_cols:
                le = LabelEncoder()
                df[col] = le.fit_transform(df[col])
        elif num == "6":
            # replace null values with 'N/A'
            df.replace(['', ' ', None, np.nan], 'N/A',
                       inplace=True, regex=True)
        elif num == "7":
            # Convert string columns to lowercase
            str_cols = df.select_dtypes(include='object').columns
            df[str_cols] = df[str_cols].apply(lambda x: x.str.lower())
        elif num == "8":
            # Sort Dataframe
            df = sort_dataframe(df, sortColumn)
        else:
            print("No Option Available")

    return df


class FilesViewSet(viewsets.ModelViewSet):
    queryset = Files.objects.all()
    serializer_class = FilesSerializer


def get_file_data(request, title, no_of_rows):

    files = Files.objects.get(id=title)

    file_path = files.pdf.path

    df = read_file(file_path)

    string_array_str = request.GET.get('stringArray')

    string_array = json.loads(string_array_str)
    
    new_list = []

    for item in string_array:
        if item not in new_list:
            new_list.append(item)

    data = []

    if no_of_rows != "All":
        filtered_df = df[new_list].head(int(no_of_rows))
    else:
        filtered_df = df[new_list]
        

    return filtered_df


@api_view(['GET'])
def upload_files_data(request, title):

    files = Files.objects.get(id=title)

    file_path = files.pdf.path

    final = read_file(file_path)

    if not file_path.split('.')[-1] == 'xlsx':

        with default_storage.open(file_path, 'r') as f:
            contents = f.read()

            lines = contents.split('\n')
            header_row = lines[0]
            column_names = header_row.split(',')

            return JsonResponse({"success": column_names})
    else:
        with default_storage.open(file_path, 'rb') as f:
            workbook = openpyxl.load_workbook(f)

            worksheet = workbook.active
            header_row = [cell.value for cell in worksheet[1]]
            return JsonResponse({"success": header_row})


@api_view(['GET'])
def upload_files_data1(request):

    return JsonResponse({"success": False})


@api_view(['GET'])
def filter_files_data(request, title):

    no_of_rows = request.GET.get('numRows')

    filtered_df = get_file_data(request, title, no_of_rows)

    filtered_data = filtered_df.to_dict(orient='records')

    for d in filtered_data:
        for key, value in d.items():
            if isinstance(value, float) and math.isnan(value):
                d[key] = None

    return JsonResponse({'result': filtered_data})


@api_view(['GET'])
def start_transformation(request, title):

    transformationOptions = request.GET.get('transformationOptions')
    sortColumn = request.GET.get('sortColumn')

    no_of_rows = request.GET.get('numRows')

    filtered_df = get_file_data(request, title, no_of_rows)
    
    filtered_df = filtered_df.drop_duplicates()


    # Tranformation Steps:

    if not len(transformationOptions) == 0:
        filtered_df = transform_dataframe(
            filtered_df, json.loads(transformationOptions), sortColumn)

    filtered_data = filtered_df.to_dict('records')

    for d in filtered_data:
        for key, value in d.items():
            if isinstance(value, float) and math.isnan(value):
                d[key] = None

    return JsonResponse({'result': filtered_data})


@api_view(['GET'])
def start_loading(request):
    string_array_str = request.GET.get('stringArray')
    tableName = request.GET.get('tableName')

    string_array = json.loads(string_array_str)

    df = pd.DataFrame(string_array)

    df = df.convert_dtypes()

    tableName = re.sub("[^0-9a-zA-Z]+", "_", tableName)

    columns, dataTypes, tableCols = get_column_name(df)

    df_data = df[columns]
    records = df_data.values.tolist()

    result = ""
    returnObj = {}
    # * Cheching if table already exists *

    tables = get_all_tables()

    if tableName in tables:
        result = "Table already Exists"
        returnObj["status"] = False
        returnObj["data"] = result

    else:
        # * Creting the Table *
        sql_createTable = create_table_query(tableName, dataTypes)

        print(sql_createTable)
        cursor = conn.cursor()

        try:
            cursor.execute(sql_createTable)
            cursor.commit()

        except Exception as e:
            cursor.rollback()
            result = str(e)

        finally:
            print("Table is Created")

        # * Inserting the Data *
        sql_insert = insert_sql_query(tableName, tableCols)
        try:
            for singleRecord in records:
                if pd.isna(singleRecord).sum() > 0:
                    new_cols, new_rec = get_non_null_cols_and_record(
                        tableCols, singleRecord
                    )
                    new_sql_insert = insert_sql_query(tableName, new_cols)
                    cursor.execute(new_sql_insert, new_rec)

                else:
                    cursor.execute(sql_insert, singleRecord)
                    cursor.commit()

            result = "Data Has been Loaded"

        except Exception as e:
            cursor.rollback()
            result = str(e)

        finally:
            if result == "Data Has been Loaded":
                returnObj["status"] = True
            else:
                returnObj["status"] = False

            returnObj["data"] = result
            cursor.close()

    return HttpResponse(json.dumps(returnObj), content_type="application/json")


def register_user(request):

    username = request.GET.get('username')
    email = request.GET.get('email')
    password = request.GET.get('password')
    cpassword = request.GET.get('cpassword')

    result = ""
    returnObj = {}

    cursor = conn.cursor()

    sql_data = f'''SELECT username,email FROM RegisterTable'''
    cursor.execute(sql_data)

    data = cursor.fetchall()

    usernames = []
    emails = []
    for d in data:
        i = 0
        for item in d:
            if i == 0:
                usernames.append(item)
            else:
                emails.append(item)
            i += 1

    if username in usernames:
        returnObj['data'] = "Username Already Exists"
        returnObj['status'] = False
    elif email in emails:
        returnObj['data'] = "Email Already Exists"
        returnObj['status'] = False
    elif password != cpassword:
        returnObj['data'] = "Password not matched"
        returnObj['status'] = False
    else:
        sql_insert = f'''INSERT INTO RegisterTable VALUES('{username}','{email}','{password}','{cpassword}');'''
        cursor.execute(sql_insert)
        cursor.commit()

        User = get_user_model()
        User.objects.create(username=username, email=email, password=password)

        returnObj['data'] = "Registered successfully"
        returnObj['status'] = True

    return HttpResponse(json.dumps(returnObj), content_type="application/json")


def login_user(request):

    username = request.GET.get('email')
    password = request.GET.get('password')

    returnObj = {'status': False, 'data': "Not matched"}

    user = authenticate(request=request, username=username, password=password)
    if user is not None:
        login(request, user)
        returnObj["data"] = "Login Successfull"
        returnObj["status"] = True

    return HttpResponse(json.dumps(returnObj), content_type="application/json")


def logout_user(request):
    logout(request)
    return HttpResponse({"data": "Logged Out"}, content_type="application/json")
